import shutil
import pytest
import pymupdf
from pptx import Presentation
from lecture_study.slides import render_deck


def test_pdf_renders_every_page_including_image_only_pages(tmp_path):
    path = tmp_path / "slides.pdf"
    with pymupdf.open() as document:
        page = document.new_page(width=900, height=500)
        page.insert_text((40, 40), "Gradient descent: w = w - alpha * gradient")
        document.new_page(width=900, height=500).draw_rect(pymupdf.Rect(20, 20, 300, 300), color=(1, 0, 0))
        document.save(path)
    pages = render_deck(path, tmp_path / "output")
    assert len(pages) == 2
    assert "Gradient descent" in pages[0]["text"]
    assert pages[1]["image"].stat().st_size > 0
    assert max(pymupdf.Pixmap(pages[0]["image"]).width, pymupdf.Pixmap(pages[0]["image"]).height) <= 1281


def test_rejects_too_many_or_encrypted_pages(tmp_path):
    path = tmp_path / "slides.pdf"
    with pymupdf.open() as document:
        for _ in range(121): document.new_page()
        document.save(path)
    with pytest.raises(ValueError, match="120"):
        render_deck(path, tmp_path / "output")
    private = tmp_path / "encrypted.pdf"
    with pymupdf.open() as document:
        document.new_page()
        document.save(private, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="user")
    with pytest.raises(ValueError, match="Password"):
        render_deck(private, tmp_path / "private")


@pytest.mark.skipif(not shutil.which("libreoffice"), reason="PPTX conversion is checked in the runtime image with LibreOffice")
def test_pptx_renders_pages_and_preserves_speaker_notes(tmp_path):
    source = tmp_path / "slides.pptx"
    presentation = Presentation()
    for title in ("Gradient descent", "Regularization"):
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = title
        slide.notes_slide.notes_text_frame.text = "Teacher notes, not recorded speech"
    presentation.save(source)
    pages = render_deck(source, tmp_path / "render")
    assert len(pages) == 2
    assert "Teacher notes" in pages[0]["speakerNotes"]
    assert "Gradient descent" in pages[0]["text"]
